"""Build the companion Excel workbook from model.py + evidence.py."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from model import N, R, PAPERS
from evidence import OBS, SEARCHES

NODES = {n["id"]: n for n in N}
HDR_FILL = PatternFill("solid", fgColor="DDE6EF")
HDR_FONT = Font(bold=True)

wb = Workbook()


def sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- Summary
ws = wb.active
ws.title = "Summary"
summary = [
    ("Title", "NRF2 signalling in colorectal cancer malignant progression - baseline map"),
    ("Built", "2026-10-29"),
    ("Central factor", "NRF2 / NFE2L2"),
    ("Disease / tissue", "Colorectal cancer (human)"),
    ("Species", "Homo sapiens (main pathway). Non-human / non-CRC systems appear only as candidate pathway."),
    ("Figure question",
     "What raises or lowers NRF2 activity in CRC cells, what assists each step, and which malignant phenotypes follow?"),
    ("Detail level", "Standard (28 nodes, 33 reactions, 2 helper edges)"),
    ("Mode", "Baseline map. No drug-intervention layer overlaid."),
    ("", ""),
    ("ASSUMPTIONS", ""),
    ("A1", "Cell lines / subtypes / mutations were not specified by the requester; whatever each source used was recorded as-is in RawExtraction."),
    ("A2", "No stimulus or culture condition was prioritised; perturbations are recorded per observation."),
    ("A3", "All evidence in this build comes from PubMed abstracts. No figure panel was opened. Every row is marked confirm=abstract."),
    ("A4", "mRNA nodes are gene-product-level functional summaries. They are NOT a claim that transcription is the dominant mechanism."),
    ("A5", "The cytosol/nucleus split of NRF2 and the cNRF2-PSMD4 branch were deliberately excluded at the requester's instruction."),
    ("", ""),
    ("MAIN CONCLUSIONS", ""),
    ("C1", "In CRC, NRF2 protein level/activity is raised by several converging post-transcriptional routes - GDF15-PI3K/AKT/GSK3b stabilisation, UCKL1 stabilisation, HNRNPC mRNA+protein stabilisation, and CHTOP/SENP3 deSUMOylation - rather than by a single transcriptional input."),
    ("C2", "The best-supported effector arm is NRF2 -> SLC7A11 -> ferroptosis resistance -> chemoresistance, with TXNRD1 and HMOX1 as parallel chemoresistance arms."),
    ("C3", "FOXO3 -> KEAP1 is the one negative-regulatory input retrieved with direct promoter-binding evidence."),
    ("C4", "GDF15 and NRF2 form a documented positive feedback loop (ChIP + luciferase)."),
    ("", ""),
    ("REMAINING GAPS", ""),
    ("G1", "Anoikis resistance: no CRC evidence retrieved. Drawn as candidate from a fibrosarcoma/lung study (C3)."),
    ("G2", "Stemness / CSC growth: no CRC paper perturbing NRF2 and measuring self-renewal retrieved. Candidate only."),
    ("G3", "GPX4: regulation by NRF2 in CRC is concurrent/association only. Labelled Mediator unresolved."),
    ("G4", "ABCB1/MDR1: association only; MDR1 was the engineered driver in the source model."),
    ("G5", "The ubiquitin ligase downstream of GSK3b (e.g. bTrCP) is not drawn - Mediator unresolved."),
    ("G6", "Conflict recorded: one source reports NRF2 only weakly induced by oxaliplatin, with EMT attributed to ROS-AKT-Snail (see Review R-CONFLICT-1)."),
    ("", ""),
    ("LEGEND", ""),
    ("Arrowhead", "Type of action: mim-conversion / mim-catalysis / mim-stimulation / mim-inhibition / mim-transcription-translation"),
    ("Line style", "Solid = main pathway (C1-C2 context fit, E1-E3 evidence). Broken = candidate pathway (other system, or association only)."),
    ("Boxes", "Subcellular compartment. The nucleus is drawn twice; both regions are the same compartment (marked on the figure)."),
    ("State mark", "'S' on NRF2 = SUMO conjugation, removed by SENP3."),
    ("Colour", "Carries no meaning in this map."),
    ("", ""),
    ("EVIDENCE LEVEL (directness)", ""),
    ("E1", "Direct: purified/in vitro reaction, direct binding, ChIP + reporter mutation"),
    ("E2", "Perturbation-dependent: knockdown/knockout/inhibitor/overexpression changes the downstream readout"),
    ("E3", "Indirect: perturbation effect shown but a mediator may intervene"),
    ("E4", "Association only: correlation, co-expression, co-localisation"),
    ("", ""),
    ("CONTEXT FIT", ""),
    ("C1", "Species, disease, cell and condition all match the target"),
    ("C2", "Species matches; cell / disease / condition partly differ"),
    ("C3", "Different species, or unrelated cell line"),
    ("C4", "Cell-free, computational prediction, or database annotation only"),
    ("", ""),
    ("VERIFICATION STATUS", ""),
    ("Format", "GPML validated mechanically (well-formedness, GraphId uniqueness, GraphRef/Anchor resolution, node overlap by coordinate). See the delivery report."),
    ("Rendering", "NOT opened in PathVisio itself. 'Format-validated' and 'rendering-confirmed on the real application' are different claims; only the former applies."),
]
for k, v in summary:
    ws.append([k, v])
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 120
for row in ws.iter_rows(min_row=1):
    row[0].font = Font(bold=True)
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# ---------------------------------------------------------------- RawExtraction
sheet("RawExtraction",
      ["ObsID", "PaperID", "Factor", "Biological kind", "Upstream", "Downstream",
       "Measured quantity", "Observed change", "Direction", "Species", "Cell / tissue",
       "Disease / mutation", "Stimulus / intervention", "Dose", "Time", "Method",
       "Figure / panel", "Evidence level", "Context fit", "Confirmation",
       "Unverified points / limits"],
      [[o["id"], o["paper"], o["factor"], o["kind"], o["upstream"], o["downstream"],
        o["measured"], o["change"], o["direction"], o["species"], o["cell"], o["disease"],
        o["stimulus"], o["dose"], o["time"], o["method"], o["figure"], o["E"], o["C"],
        o["confirm"], o["limits"]] for o in OBS],
      [9, 9, 20, 14, 18, 24, 18, 46, 16, 16, 26, 20, 30, 8, 8, 30, 13, 13, 11, 12, 48])

# ---------------------------------------------------------------- Review
REVIEW = [
    ("RV01", "Aggregation", "NFE2L2 mRNA + NRF2 protein",
     "Kept BOTH nodes. mRNA retained because HNRNPC acts on transcript stability; protein retained because SUMOylation, KEAP1-dependent degradation and TF activity all operate at protein level.",
     "Not aggregated to a single mRNA node: the protein-level mechanisms are the point of the map."),
    ("RV02", "State split", "NRF2 (SUMOylated) -> NRF2",
     "Split kept. deSUMOylation by SENP3 is the step that sustains HMOX1 expression, so the two states have different downstream consequence.",
     "No separate 'active' node added - the state mark plus the comment carries the activity meaning."),
    ("RV03", "State split", "NRF2 cytosol vs nuclear",
     "NOT split. Requester excluded the cNRF2/PSMD4 branch, and without it the localisation distinction would not change any drawn downstream action.",
     "The supporting paper (PMID 29899863) is recorded in Papers as held back, not deleted."),
    ("RV04", "Chain", "gene -> mRNA -> protein",
     "Not auto-added. mRNA->protein edges exist only where a source actually reasons across both levels (NFE2L2, KEAP1, SLC7A11, GDF15), and each is tagged as a gene-product-level summary.",
     "ABCB1, TXNRD1, HMOX1 are left at mRNA level only; their protein nodes were not invented."),
    ("RV05", "Isoform", "AKT",
     "Source says 'canonical PI3K/AKT/GSK3b'. Drawn as 'AKT1 and/or AKT2 and/or AKT3'. Xref points to AKT1 as a placeholder only.",
     "Not assigned to a single isoform."),
    ("RV06", "Helper typing", "SENP3 vs CHTOP on R006",
     "SENP3 drawn as mim-catalysis (enzyme). CHTOP drawn as mim-stimulation (recruitment/scaffold). Both attach to the Anchor on the conversion edge, not to the product node.",
     "Distinguishing enzyme from scaffold was required by the brief."),
    ("RV07", "Downgrade", "NRF2 -> GPX4 (R024)",
     "Downgraded to candidate (broken, E4). The review P001 asserts direct transcriptional control, but a review cannot be the sole basis; the retrieved originals show only concurrent movement or pharmacological co-modulation.",
     "Labelled Mediator unresolved."),
    ("RV08", "Downgrade", "NRF2 -> ABCB1 (R018) and ABCB1 -> chemoresistance (R027)",
     "Drawn as candidate (broken, E4). In the source, MDR1 was the engineered driver of the resistant line, so NRF2-dependence of ABCB1 is not established.",
     "Kept rather than dropped because the resistance phenotype itself is NRF2-dependent (ML-385 data, OBS012)."),
    ("RV09", "Candidate", "Anoikis arm (R031, R032, N020, N034)",
     "Retained as candidate at requester's instruction. Evidence is fibrosarcoma / lung adenocarcinoma (C3).",
     "Three separate CRC-targeted queries returned 0-1 hits (S003). Gap reported, not filled."),
    ("RV10", "Candidate", "Stemness arm (R033, N035)",
     "Retained as candidate at requester's instruction. Source describes NRF2/KEAP1 downstream of a drug combination without NRF2 perturbation.",
     "Three queries returned zero (S004). Gap reported, not filled."),
    ("RV11", "Exclusion", "PMID 30993753 (Nrf2/STAT3, HT-29)",
     "EXCLUDED. Flagged as Retracted Publication in PubMed.",
     "Not used for any node, edge or claim. Retained in Papers with the exclusion reason."),
    ("RV12", "Held back", "PMID 26733168 (oxaliplatin ROS-AKT-Snail EMT)",
     "HELD BACK from the figure but recorded. It reports NRF2 only weakly induced by oxaliplatin and attributes EMT to a ROS-AKT-Snail axis.",
     "See R-CONFLICT-1 below."),
    ("R-CONFLICT-1", "Conflict", "NRF2 -> migration & invasion (R030)",
     "Two findings disagree. OBS010 (P007) shows shNrf2 / brusatol enhance, and tBHQ attenuates, inhibition of migration and invasion - i.e. NRF2 supports invasion. OBS016 (P901) finds NRF2 barely induced by oxaliplatin and assigns EMT to ROS-AKT-Snail.",
     "ADOPTED OBS010 for the drawn edge, because it perturbs NRF2 directly in both directions; OBS016 tests a drug stimulus rather than NRF2 itself. Node N033 carries a comment pointing here."),
    ("RV13", "Non-monotonic", "CHTOP (OBS006)",
     "The source reports that BOTH restoring and further depleting CHTOP resensitised 5-FU-resistant cells.",
     "The figure draws only the deSUMOylation step, which is what the mechanism section supports. The non-monotonic phenotype relation is recorded here and not drawn as a simple arrow."),
    ("RV14", "Display repeat", "Nucleus drawn twice",
     "The nucleus appears as two regions so that upstream transcription stays above NRF2 and target-gene transcription stays below it, preserving top-to-bottom flow.",
     "Both regions are labelled as the same compartment to prevent misreading them as different organelles."),
]
sheet("Review", ["ReviewID", "Decision type", "Target", "Judgement", "Consequence / note"],
      [list(r) for r in REVIEW], [14, 16, 42, 86, 72])

# ---------------------------------------------------------------- Reactions
sheet("Reactions",
      ["ReactionID (GraphId)", "From (state before)", "From ID", "To (state after)", "To ID",
       "Reaction type", "Arrowhead", "Helper factor(s) + role", "Pathway layer",
       "Evidence level", "Context fit", "Cross-paper chain", "Papers", "Note"],
      [[r["id"], NODES[r["src"]]["label"], r["src"], NODES[r["dst"]]["label"], r["dst"],
        r["kind"], r["arrow"],
        "; ".join(f"{NODES[h.split(':')[0]]['label']} ({'catalysis' if h.split(':')[1]=='mim-catalysis' else 'recruitment/scaffold'})"
                  for h in r["helper"].split(",")) if r["helper"] else "",
        r["layer"], r["E"], r["C"], "YES" if r["cross"] else "", r["papers"], r["note"]]
       for r in R],
      [20, 26, 9, 26, 9, 26, 26, 34, 12, 12, 10, 14, 14, 70])

# ---------------------------------------------------------------- CandidateSearch
sheet("CandidateSearch",
      ["SearchID", "Missing mechanism", "Database", "Query", "Search date", "Filters",
       "Hits", "Candidate factor(s)", "Ranking reason", "Experimental system",
       "Decision", "Note"],
      [[s["id"], s["gap"], s["db"], s["query"], s["date"], s["filters"], str(s["hits"]),
        s["candidates"], s["rank_reason"], s["system"], s["decision"], s["note"]]
       for s in SEARCHES],
      [10, 46, 11, 60, 12, 16, 10, 44, 70, 34, 54, 58])

# ---------------------------------------------------------------- Entities
sheet("Entities",
      ["NodeID (GraphId)", "Label", "Official gene symbol", "Biological type",
       "Modification / localisation / activity state", "Compartment (as drawn)",
       "Pathway layer", "Xref database", "Xref ID", "Note"],
      [[n["id"], n["label"],
        n["label"].replace(" mRNA", "").split(" (")[0] if n["type"] != "Pathway" else "",
        n["type"],
        "SUMOylated" if n["state"] else ("phenotype node" if n["type"] == "Pathway" else "none"),
        n["comp"], n["layer"], n["xdb"], n["xid"], n["note"]] for n in N],
      [18, 34, 20, 14, 30, 16, 12, 14, 10, 74])

# ---------------------------------------------------------------- Papers
sheet("Papers",
      ["PaperID", "First author", "Year", "Journal", "Title", "PMID", "DOI", "URL",
       "Type", "Figure / text confirmation", "Note"],
      [[p["id"], p["first"], p["year"], p["journal"], p["title"], p["pmid"], p["doi"],
        f"https://doi.org/{p['doi']}", p["type"], p["confirm"], p["note"]] for p in PAPERS],
      [9, 24, 7, 28, 78, 11, 32, 40, 11, 24, 62])

wb.save("NRF2_CRC_baseline_evidence.xlsx")
print("wrote NRF2_CRC_baseline_evidence.xlsx")
print("sheets:", wb.sheetnames)
